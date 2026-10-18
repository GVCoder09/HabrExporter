import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from lxml import html as lhtml

from conftest import AVATAR_AUTHOR, COMPANY_LOGO, IMG_BROKEN, IMG_FIGURE, IMG_FORMULA, IMG_LAZY
from habr_exporter.exporters import (
    DocxExporter,
    EpubExporter,
    ExportFormat,
    ExportOptions,
    Fb2Exporter,
    HtmlExporter,
    PdfExporter,
    collect_image_urls,
    get_exporter,
)
from habr_exporter.exporters.html import make_image_src as make_html_image_src
from habr_exporter.exporters.pdf import PdfExportError, PdfPrinter, html_to_pdf, make_image_src
from habr_exporter.images import ImageMode, ImageStore


@pytest.fixture
def images(client, article):
    store = ImageStore(client)
    store.fetch_all(collect_image_urls(article))
    return store


def test_get_exporter():
    assert isinstance(get_exporter(ExportFormat.PDF), PdfExporter)
    assert isinstance(get_exporter(ExportFormat.HTML), HtmlExporter)
    assert isinstance(get_exporter(ExportFormat.DOCX), DocxExporter)
    assert isinstance(get_exporter(ExportFormat.EPUB), EpubExporter)
    assert isinstance(get_exporter(ExportFormat.FB2), Fb2Exporter)
    assert ExportFormat.DOCX.extension == ".docx"


def test_format_order_and_beta_labels():
    assert list(ExportFormat) == [
        ExportFormat.PDF,
        ExportFormat.HTML,
        ExportFormat.DOCX,
        ExportFormat.EPUB,
        ExportFormat.FB2,
        ExportFormat.MARKDOWN,
        ExportFormat.JSON,
    ]
    labels = ["PDF", "HTML", "DOCX (beta)", "EPUB (beta)", "FB2 (beta)", "MD (beta)", "JSON"]
    assert [f.label for f in ExportFormat] == labels
    assert not ExportFormat.PDF.beta
    assert not ExportFormat.HTML.beta
    assert not ExportFormat.JSON.beta
    assert [f.extension for f in ExportFormat] == [
        ".pdf",
        ".html",
        ".docx",
        ".epub",
        ".fb2",
        ".md",
        ".json",
    ]
    assert [f for f in ExportFormat if f.needs_qt] == [
        ExportFormat.PDF,
        ExportFormat.DOCX,
        ExportFormat.FB2,
    ]


def test_collect_image_urls(article):
    urls = collect_image_urls(article)
    assert urls[0] == IMG_FIGURE
    assert {IMG_LAZY, IMG_FORMULA, IMG_BROKEN, AVATAR_AUTHOR, COMPANY_LOGO} <= set(urls)
    assert len(urls) == len(set(urls))


# --- HTML ---------------------------------------------------------------------------


def test_html_embed_is_self_contained(tmp_path, article, images):
    path = HtmlExporter().export(
        article, tmp_path / "a.html", ExportOptions(ExportFormat.HTML), images
    )
    html = path.read_text(encoding="utf-8")
    doc = lhtml.document_fromstring(html)
    assert doc.findtext(".//title") == f"{article.title} / Хабр"
    assert doc.xpath("//style")
    assert not doc.xpath("//link[@rel='stylesheet'] | //script")
    srcs = doc.xpath("//div[@class='article-body']//img/@src")
    assert any(s.startswith("data:image/png;base64,") for s in srcs)
    assert any(s.startswith("data:image/webp;base64,") for s in srcs)
    assert any(s.startswith("data:image/svg+xml;base64,") for s in srcs)
    assert IMG_BROKEN in srcs  # failed to download, so it stays a link to the original
    assert IMG_FIGURE not in srcs
    assert doc.xpath("//div[@class='article-meta-top']/img/@src")[0].startswith("data:image/png")
    # Spoilers are collapsed, as on the site.
    assert doc.xpath("//details")[0].get("open") is None


def test_html_link_mode_keeps_original_urls(tmp_path, article, images):
    options = ExportOptions(ExportFormat.HTML, ImageMode.LINK, include_comments=False)
    path = HtmlExporter().export(article, tmp_path / "a.html", options, images)
    doc = lhtml.document_fromstring(path.read_text(encoding="utf-8"))
    srcs = doc.xpath("//div[@class='article-body']//img/@src")
    assert IMG_FIGURE in srcs
    assert not any(s.startswith("data:") for s in srcs)
    assert not doc.xpath("//section[contains(@class,'comments')]")


def test_html_embed_without_store_falls_back_to_urls():
    resolve = make_html_image_src(ImageMode.EMBED, None)
    assert resolve(IMG_FIGURE) == IMG_FIGURE


# --- DOCX ---------------------------------------------------------------------------


def _all_text(doc):
    """Return the document text in order, including table cells."""
    parts = []
    for block in doc.element.body.iterchildren():
        if block.tag == qn("w:p"):
            parts.append("".join(t.text or "" for t in block.iter(qn("w:t"))))
        elif block.tag == qn("w:tbl"):
            parts.extend(
                "".join(t.text or "" for t in cell.iter(qn("w:t")))
                for cell in block.iter(qn("w:tc"))
            )
    return "\n".join(parts)


def _hyperlinks(path):
    with zipfile.ZipFile(path) as zf:
        rels = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
    return {
        r.get("Target")
        for r in rels
        if r.get("TargetMode") == "External" and r.get("Type").endswith("/hyperlink")
    }


def _document_xml(path):
    with zipfile.ZipFile(path) as zf:
        return etree.fromstring(zf.read("word/document.xml"))


# Child element order per the OOXML schema: Word refuses to open files that violate it.
_PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]  # fmt: skip
_RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
    "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color",
    "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
    "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]  # fmt: skip
_BDR_ORDER = ["top", "left", "bottom", "right", "between", "bar"]


def _assert_order(parent, order):
    names = [etree.QName(child).localname for child in parent]
    positions = [order.index(n) for n in names]
    assert positions == sorted(positions), names


def test_docx_schema_order(tmp_path, article, images, qapp):
    path = DocxExporter().export(
        article, tmp_path / "a.docx", ExportOptions(ExportFormat.DOCX), images
    )
    root = _document_xml(path)
    for ppr in root.iter(qn("w:pPr")):
        _assert_order(ppr, _PPR_ORDER)
    for rpr in root.iter(qn("w:rPr")):
        if rpr.getparent().tag == qn("w:r"):
            _assert_order(rpr, _RPR_ORDER)
    for bdr in root.iter(qn("w:pBdr")):
        _assert_order(bdr, _BDR_ORDER)


def test_docx_content(tmp_path, article, images, qapp):
    path = DocxExporter().export(
        article, tmp_path / "a.docx", ExportOptions(ExportFormat.DOCX), images
    )
    doc = Document(str(path))
    text = _all_text(doc)
    assert doc.core_properties.title == article.title
    assert doc.paragraphs[1].text == article.title
    # Order: body -> tags -> company -> author -> comments.
    marks = (
        "Вступление",
        "Теги:",
        "Делаем всё на свете",
        "Инженер-тестировщик",
        "Первый комментарий",
    )
    order = [text.index(s) for s in marks]
    assert order == sorted(order)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings[:2] == ["Раздел первый", "Подраздел"]
    assert "•\tПервый пункт" in text
    assert "◦\tВложенный пункт" in text
    assert "3.\tТретий" in text
    assert "4.\tЧетвёртый" in text
    assert "▸ Под спойлером" in text
    assert "Скрытый текст." in text
    assert "НЛО прилетело" in text
    assert "Подпись к рисунку" in text
    assert "строка после переноса" in text
    assert "объединённая ячейка" in text

    code = next(p for p in doc.paragraphs if "def greet" in p.text)
    assert code.runs[0].font.name == "Consolas"
    assert len({r.font.color.rgb for r in code.runs if r.font.color.rgb}) > 1
    assert "\n" in code.text or code._p.xpath(".//w:br")

    links = _hyperlinks(path)
    assert {"https://example.com/", "https://habr.com/ru/articles/1/", "https://t.me/acme"} <= links
    assert IMG_BROKEN in links  # a failed image becomes a link
    assert "https://embedd.srv.habr.com/iframe/abc123" in text

    pictures = doc.inline_shapes
    # figure, webp, formula (svg), image in a comment, avatars and logo
    assert len(pictures) >= 6
    assert (
        max(p.width for p in pictures) <= doc.sections[0].page_width - doc.sections[0].left_margin
    )

    table = next(t for t in doc.tables if t.style is not None and t.style.name == "Table Grid")
    assert table.cell(0, 0).text == "Ключ"
    assert table.cell(2, 0).text == table.cell(2, 1).text == "объединённая ячейка"


def test_docx_comment_indentation(tmp_path, article, images, qapp):
    path = DocxExporter().export(
        article, tmp_path / "a.docx", ExportOptions(ExportFormat.DOCX), images
    )
    doc = Document(str(path))

    def indent_of(substr):
        paragraph = next(p for p in doc.paragraphs if substr in p.text)
        return paragraph.paragraph_format.left_indent or 0

    root = indent_of("Первый комментарий")
    reply = indent_of("Ответ автора.")
    nested = indent_of("Согласна.")
    assert root < reply < nested


def test_docx_link_mode(tmp_path, article, qapp):
    options = ExportOptions(ExportFormat.DOCX, ImageMode.LINK, include_comments=False)
    path = DocxExporter().export(article, tmp_path / "a.docx", options, None)
    doc = Document(str(path))
    assert len(doc.inline_shapes) == 0
    assert IMG_FIGURE in _hyperlinks(path)
    assert "🖼" in _all_text(doc)
    assert "Первый комментарий" not in _all_text(doc)


def test_docx_link_mode_links_pictures_to_habr(tmp_path, article, images, qapp):
    options = ExportOptions(ExportFormat.DOCX, ImageMode.LINK, include_comments=False)
    assert options.downloads_images  # sizes are needed even when pictures are only linked
    path = DocxExporter().export(article, tmp_path / "a.docx", options, images)
    with zipfile.ZipFile(path) as z:
        document = z.read("word/document.xml").decode("utf-8")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "r:link=" in document
    assert f'Target="{IMG_FIGURE}" TargetMode="External"' in rels
    assert IMG_FIGURE not in _hyperlinks(path)


def test_docx_without_comments_list(tmp_path, article, qapp):
    article.comments = []
    article.company = None
    path = DocxExporter().export(
        article, tmp_path / "a.docx", ExportOptions(ExportFormat.DOCX, ImageMode.LINK), None
    )
    assert "Комментариев пока нет" in _all_text(Document(str(path)))


# --- PDF ----------------------------------------------------------------------------


def test_make_image_src(images):
    embed = make_image_src(ImageMode.EMBED, images)
    assert embed(IMG_FIGURE).startswith("data:image/png;base64,")
    assert embed(IMG_BROKEN) is None
    assert (
        make_image_src(ImageMode.LINK, images)(IMG_FIGURE) == IMG_FIGURE
    )  # Chromium loads it while printing
    assert make_image_src(ImageMode.LINK, None)(IMG_FIGURE) == IMG_FIGURE
    assert make_image_src(ImageMode.EMBED, None)(IMG_FIGURE) is None


def test_pdf_render_expands_spoilers(article, images):
    html = PdfExporter().render(article, ExportOptions(ExportFormat.PDF), images)
    assert '<details class="spoiler" open' in html
    assert "data:image/png;base64," in html


def test_pdf_export(tmp_path, article, images, qapp):
    path = PdfExporter().export(
        article, tmp_path / "a.pdf", ExportOptions(ExportFormat.PDF), images
    )
    data = path.read_bytes()
    assert data.startswith(b"%PDF")
    assert len(data) > 10_000


def test_pdf_export_link_mode(tmp_path, article, qapp):
    options = ExportOptions(ExportFormat.PDF, ImageMode.LINK)
    path = PdfExporter().export(article, tmp_path / "b.pdf", options, None)
    assert path.read_bytes().startswith(b"%PDF")


def test_pdf_bad_target_reports_error(tmp_path, qapp):
    with pytest.raises(PdfExportError):
        html_to_pdf("<p>x</p>", tmp_path / "missing-dir" / "x.pdf")


def test_pdf_printer_timeout(tmp_path, qapp, qtbot):
    printer = PdfPrinter(timeout_ms=1)
    with qtbot.waitSignal(printer.failed, timeout=10_000) as blocker:
        printer.print_html("<p>x</p>", tmp_path / "x.pdf")
    assert "время ожидания" in blocker.args[0]


def test_pdf_printer_requires_main_thread(qapp):
    from PyQt6.QtCore import QThread

    errors = []

    class Worker(QThread):
        def run(self):
            try:
                PdfPrinter()
            except PdfExportError as exc:
                errors.append(str(exc))

    worker = Worker()
    worker.start()
    worker.wait(5000)
    assert errors
    assert "главного потока" in errors[0]
